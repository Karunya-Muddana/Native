# app/tools/excel_info.py
from pathlib import Path
from langchain_core.tools import tool
from openpyxl import load_workbook

from app.tools.paths import BASE_DIR, INPUT_DIR, OUTPUT_DIR, resolve_file as _resolve


@tool
def get_excel_info(filename: str, location: str = "input") -> str:
    """
    Get basic information about an Excel file: sheet names, and each sheet's
    dimensions (rows x columns). Call this BEFORE read_excel_range on an
    unfamiliar spreadsheet, so you know which sheet to read and how big it is.
    location is 'input' for a user-uploaded file, or 'output' for one saved by
    a previous tool.
    """
    try:
        file_path = _resolve(filename, location)
        if isinstance(file_path, str):
            return file_path

        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        lines = [f"File: {filename}", f"Sheets: {wb.sheetnames}"]
        for name in wb.sheetnames:
            ws = wb[name]
            lines.append(f"  '{name}': {ws.max_row} rows x {ws.max_column} columns")
        wb.close()
        return "\n".join(lines)
    except Exception as e:
        return f"Error: {str(e)}"


@tool
def read_excel_range(
    filename: str,
    sheet_name: str,
    start_row: int,
    end_row: int,
    location: str = "input",
) -> str:
    """
    Read a range of rows from one sheet in an Excel file, returned as a simple
    text table. Call get_excel_info first to see valid sheet names and row
    counts. Rows are 1-indexed (matching Excel's own row numbers). Use this to
    read a large spreadsheet in chunks across multiple calls, rather than
    trying to read it all at once.
    location is 'input' for a user-uploaded file, or 'output' for one saved by
    a previous tool.
    """
    try:
        file_path = _resolve(filename, location)
        if isinstance(file_path, str):
            return file_path

        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            return f"Error: sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}"

        ws = wb[sheet_name]
        total_rows = ws.max_row
        if start_row > total_rows:
            return f"Error: start_row {start_row} is beyond the sheet ({total_rows} rows total)."

        end_row = min(end_row, total_rows)
        lines = []
        for row in ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True):
            lines.append(" | ".join("" if v is None else str(v) for v in row))
        wb.close()

        table = "\n".join(lines)
        if end_row >= total_rows:
            note = f"\n\n[End of sheet reached — {total_rows} rows total.]"
        else:
            note = f"\n\n[Showing rows {start_row}-{end_row} of {total_rows} total.]"

        return table + note
    except Exception as e:
        return f"Error: {str(e)}"