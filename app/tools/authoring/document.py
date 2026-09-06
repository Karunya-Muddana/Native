# app/tools/authoring/document.py
"""`write_document` — a .docx the operator can sign, file or send.

The flagship workflow of this project ends in a document, not a chat message:
scanned inspection report → OCR → thickness data → knowledge base rules →
calculation → an approval note someone puts their name on. Everything before
this tool produces evidence; this is the step that produces the deliverable.

The file is re-opened and read back after writing, and the tool reports what is
actually in it. That is not a formality — the agent's own summary of a file it
generated is a guess about what it *meant* to write, and the two diverge.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from langchain_core.tools import tool

from app.tools.authoring.blocks import (
    SpecError,
    image_path,
    output_path,
    parse_blocks,
    size_note,
)

MAX_IMAGE_WIDTH = Inches(6.0)  # fits inside the default page margins
TABLE_STYLE = "Table Grid"


@tool
def write_document(filename: str, title: str, content: str, subtitle: str = "") -> str:
    """
    Write a Word document (.docx) to /sandbox/output/ — an inspection note, an
    approval memo, a findings report, anything the user will file, sign or send.

    title is the document's heading. subtitle is optional, and is the place for
    a reference line: vessel tag, report number, date, author.

    content is a JSON list of blocks, in the order they should appear:
      {"type": "heading",   "text": "Findings", "level": 2}
      {"type": "paragraph", "text": "Measured minimum thickness is 10.9 mm."}
      {"type": "bullets",   "items": ["Point P4 below limit", "Re-inspect in 6 months"]}
      {"type": "table",     "header": ["Point", "mm"], "rows": [["P1", 12.4], ["P4", 10.9]]}
      {"type": "image",     "filename": "thickness_chart.png", "caption": "Fig 1"}
      {"type": "pagebreak"}

    An image block embeds a file already in /sandbox/output/ or /sandbox/input/ —
    generate charts with python_runner first, then reference them by name here.

    Use this when the user asks for a document, note, memo, report or letter.
    For tabular data they will sort or filter, use create_spreadsheet; for
    something they will present, use create_presentation. The tool reports what
    the finished file actually contains — quote those numbers in your reply
    rather than restating what you intended to write.
    """
    try:
        blocks = parse_blocks(content)
        path = output_path(filename, ".docx")
    except SpecError as e:
        return f"Error: {e}"

    document = Document()
    _style(document)

    title = (title or "").strip()
    if title:
        document.add_heading(title, level=0)
    subtitle = (subtitle or "").strip()
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.italic = True

    try:
        for block in blocks:
            _render(document, block)
    except SpecError as e:
        return f"Error: {e} Nothing was written."

    try:
        document.save(str(path))
    except PermissionError:
        return (
            f"Error: {path.name} is open in another program and could not be replaced. "
            "Ask the user to close it."
        )

    return _read_back(path, title)


def _style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)


def _render(document: Document, block: dict) -> None:
    kind = block["type"]

    if kind == "heading":
        document.add_heading(block["text"], level=block["level"])

    elif kind == "paragraph":
        document.add_paragraph(block["text"])

    elif kind == "bullets":
        for item in block["items"]:
            document.add_paragraph(item, style="List Bullet")

    elif kind == "table":
        _table(document, block)

    elif kind == "image":
        path = image_path(block)
        document.add_picture(str(path), width=MAX_IMAGE_WIDTH)
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if block["caption"]:
            caption = document.add_paragraph(block["caption"])
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs:
                run.italic = True
                run.font.size = Pt(9)

    elif kind == "pagebreak":
        document.add_page_break()


def _table(document: Document, block: dict) -> None:
    header = block["header"]
    rows = block["rows"]
    table = document.add_table(rows=0, cols=block["width"])
    table.style = TABLE_STYLE

    if header:
        cells = table.add_row().cells
        for cell, text in zip(cells, header):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text


# ── read-back ──────────────────────────────────────────────────────────────
def _read_back(path, title: str) -> str:
    """Reopen the saved file and describe what is genuinely in it.

    Counted from the document on disk rather than from the spec that produced
    it, so a block that silently failed to render shows up here as a smaller
    number instead of being reported as written.
    """
    try:
        document = Document(str(path))
    except Exception as e:
        return f"Saved /sandbox/output/{path.name}, but it could not be reopened to verify it: {e}"

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    words = sum(len(p.split()) for p in paragraphs)
    images = len(document.inline_shapes)

    lines = [
        f"Wrote /sandbox/output/{path.name} ({size_note(path)}).",
        "Read back from the saved file:",
        f"  - {len(paragraphs)} paragraph(s), {words} words",
    ]
    for i, table in enumerate(document.tables, 1):
        lines.append(f"  - table {i}: {len(table.rows)} row(s) x {len(table.columns)} column(s)")
    if images:
        lines.append(f"  - {images} embedded image(s)")

    headings = [
        p.text.strip()
        for p in document.paragraphs
        if p.style.name.startswith("Heading") and p.text.strip()
    ]
    if headings:
        lines.append("  - headings: " + " / ".join(headings[:8]))
    if title and title not in " ".join(paragraphs):
        lines.append(f"  - note: the title '{title}' is not in the saved text")

    lines.append("Open it for the user with open_on_screen if they want to look at it.")
    return "\n".join(lines)
