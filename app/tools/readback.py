# app/tools/readback.py
"""Describe a file the agent just produced, by opening it.

The failure this exists to stop: python_runner writes a correct spreadsheet, and
the agent then summarizes it from the code it *wrote* rather than from the file
that came out. If the script filtered a row the agent forgot about, the chat
answer and the attachment disagree, and the user only finds out later.

So after a run, anything new or changed in /sandbox/output/ is opened and
described here. The numbers in the tool result then come from the artifact,
which is the thing the user will actually have.

Every reader is best-effort: an unreadable file is reported as unreadable, never
raised, because a file the agent cannot introspect is still a file it produced.
"""

from pathlib import Path

SPREADSHEET = (".xlsx", ".xlsm")
DOCUMENT = (".docx",)
DECK = (".pptx",)
IMAGE = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp")
TEXT = (".txt", ".md", ".csv", ".tsv", ".json", ".log", ".py", ".html", ".xml", ".yaml", ".yml")

PREVIEW_LINES = 3
PREVIEW_CHARS = 160
MAX_SHEETS = 4


def describe(path: Path) -> str:
    """One or more lines describing what is genuinely inside `path`."""
    try:
        size = _size(path)
    except OSError as e:
        return f"{path.name} — could not be read back ({e})"

    suffix = path.suffix.lower()
    try:
        if suffix in SPREADSHEET:
            return f"{path.name} ({size}) — {_spreadsheet(path)}"
        if suffix in DOCUMENT:
            return f"{path.name} ({size}) — {_document(path)}"
        if suffix in DECK:
            return f"{path.name} ({size}) — {_deck(path)}"
        if suffix == ".pdf":
            return f"{path.name} ({size}) — {_pdf(path)}"
        if suffix in IMAGE:
            return f"{path.name} ({size}) — {_image(path)}"
        if suffix in TEXT:
            return f"{path.name} ({size}) — {_text(path)}"
    except Exception as e:
        return f"{path.name} ({size}) — written, but could not be read back ({type(e).__name__}: {e})"

    return f"{path.name} ({size})"


def _size(path: Path) -> str:
    size = path.stat().st_size
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{max(1, size // 1024)} KB"
    return f"{size / (1024 * 1024):.1f} MB"


def _spreadsheet(path: Path) -> str:
    """Sheet shapes plus the range of each numeric column.

    The value ranges are the point: they are what a claim in the chat answer can
    be checked against without opening Excel.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    try:
        parts = []
        for name in workbook.sheetnames[:MAX_SHEETS]:
            sheet = workbook[name]
            rows = list(sheet.iter_rows(values_only=True))
            summary = f"'{name}' {sheet.max_row}x{sheet.max_column}"
            if rows:
                header = rows[0]
                if header and all(v is None or isinstance(v, str) for v in header):
                    labels = [str(v) for v in header if v is not None]
                    if labels:
                        summary += " [" + ", ".join(labels[:6]) + "]"
                    summary += _ranges(header, rows[1:])
                else:
                    summary += _ranges(None, rows)
            parts.append(summary)
        if len(workbook.sheetnames) > MAX_SHEETS:
            parts.append(f"and {len(workbook.sheetnames) - MAX_SHEETS} more sheet(s)")
        return "; ".join(parts)
    finally:
        workbook.close()


def _ranges(header, body) -> str:
    out = []
    width = max((len(r) for r in body), default=0)
    for index in range(width):
        values = [
            r[index]
            for r in body
            if index < len(r) and isinstance(r[index], (int, float)) and not isinstance(r[index], bool)
        ]
        if len(values) > 1:
            label = str(header[index]) if header and index < len(header) and header[index] else f"col {index + 1}"
            out.append(f"{label} {min(values):g}–{max(values):g}")
    return (" — " + ", ".join(out[:5])) if out else ""


def _document(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    parts = [f"{len(paragraphs)} paragraph(s), {sum(len(p.split()) for p in paragraphs)} words"]
    if document.tables:
        parts.append(
            ", ".join(f"table {len(t.rows)}x{len(t.columns)}" for t in document.tables[:3])
        )
    if document.inline_shapes:
        parts.append(f"{len(document.inline_shapes)} image(s)")
    return "; ".join(parts)


def _deck(path: Path) -> str:
    from pptx import Presentation

    deck = Presentation(str(path))
    titles = []
    for slide in deck.slides:
        title = slide.shapes.title
        titles.append(title.text.strip() if title and title.text.strip() else "(untitled)")
    listed = " / ".join(titles[:5]) + ("…" if len(titles) > 5 else "")
    return f"{len(titles)} slide(s): {listed}"


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = (reader.pages[0].extract_text() or "").strip() if reader.pages else ""
    first = " ".join(text.split())[:PREVIEW_CHARS]
    return f"{len(reader.pages)} page(s)" + (f'; page 1 starts "{first}…"' if first else "")


def _image(path: Path) -> str:
    try:
        from PIL import Image

        with Image.open(str(path)) as image:
            return f"{image.width}x{image.height} {image.format}"
    except ImportError:
        return "image"


def _text(path: Path) -> str:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    if not lines:
        return "empty"
    preview = " ⏎ ".join(line.strip()[:PREVIEW_CHARS] for line in lines[:PREVIEW_LINES] if line.strip())
    return f"{len(lines)} line(s)" + (f'; starts "{preview}…"' if preview else "")


def snapshot(directory: Path) -> dict[str, float]:
    """Filename → mtime, so a later diff can tell new and rewritten files apart."""
    if not directory.is_dir():
        return {}
    return {f.name: f.stat().st_mtime for f in directory.iterdir() if f.is_file()}


def changed_since(directory: Path, before: dict[str, float]) -> list[Path]:
    """Files that appeared or were rewritten, newest last."""
    out = []
    for f in sorted(directory.iterdir(), key=lambda p: p.stat().st_mtime) if directory.is_dir() else []:
        if not f.is_file():
            continue
        if f.name not in before or f.stat().st_mtime > before[f.name]:
            out.append(f)
    return out
